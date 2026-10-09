import csv
import re
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from .domain import HEADERS, header_key, digits, reais, money

KEYS = ['cod', 'empresa', 'cnpj', 'fgts', 'consignado', 'total', 'observacoes']

def read_table(path):
    suffix = Path(path).suffix.lower()
    if suffix == '.csv':
        with open(path, encoding='utf-8-sig', newline='') as stream:
            sample = stream.read(4096); stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=';,\t')
            except csv.Error:
                dialect = csv.excel; dialect.delimiter = ';'
            data = list(csv.reader(stream, dialect))
    elif suffix == '.xlsx':
        book = load_workbook(path, read_only=True, data_only=True)
        data = list(book.active.values); book.close()
    else:
        raise ValueError('Importe um arquivo CSV ou XLSX')
    if not data:
        raise ValueError('Arquivo vazio')
    headers = [header_key(x) for x in data[0]]
    required = [header_key(x) for x in HEADERS]
    if any(headers.count(x) != 1 for x in required):
        raise ValueError('Cabeçalhos ausentes ou duplicados. Use o template do aplicativo.')
    indexes = [headers.index(x) for x in required]
    result = []
    for number, line in enumerate(data[1:], 2):
        if not any(v is not None and str(v).strip() for v in line):
            continue
        row = {key: str(line[i] if i < len(line) and line[i] is not None else '') for key,i in zip(KEYS,indexes)}
        for key in ['fgts','consignado','total']:
            raw = line[indexes[KEYS.index(key)]] if indexes[KEYS.index(key)] < len(line) else None
            if isinstance(raw, (int,float)):
                row[key] = str(raw).replace('.', ',')
        if row['cnpj'].endswith('.0'):
            row['cnpj'] = row['cnpj'][:-2]
        row['cnpj'] = digits(row['cnpj']).zfill(14) if row['cnpj'] else ''
        if len(row['cnpj']) > 14:
            raise ValueError(f'Linha {number}: CNPJ possui mais de 14 dígitos')
        for key in ['fgts', 'consignado', 'total']:
            if row[key].strip():
                try:
                    row[key] = reais(money(row[key]))
                except ValueError as exc:
                    raise ValueError(f'Linha {number}, {key}: valor monetário inválido') from exc
        row['selected'] = True
        result.append(row)
    return result

def write_table(path, rows):
    if Path(path).suffix.lower() not in {'.csv', '.xlsx'}:
        raise ValueError('Escolha a extensão .csv ou .xlsx')
    values = [[row.get(key, '') for key in KEYS] for row in rows]
    if Path(path).suffix.lower() == '.xlsx':
        book = Workbook(); sheet = book.active; sheet.title = 'Empresas'
        sheet.append(HEADERS)
        for value in values:
            sheet.append(value)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 24
        sheet.freeze_panes = 'A2'
        for cell in sheet['C']: cell.number_format = '@'
        book.save(path)
    else:
        with open(path, 'w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream, delimiter=';'); writer.writerow(HEADERS)
            # Prevent formulas when CSV is opened in spreadsheet applications.
            writer.writerows([["'"+str(v) if str(v).startswith(('=','+','-','@')) else v for v in row] for row in values])

def guide_due(path):
    text = '\n'.join(page.extract_text() or '' for page in PdfReader(path).pages)
    found = re.search(r'Pagar este documento até\s*(\d{2}/\d{2}/\d{4})', text)
    if not found:
        raise ValueError('Vencimento da guia existente não identificado no PDF')
    return found.group(1)

def validate_pdf(path, company, initial, final, due, guide):
    if not re.fullmatch(r'\d{10,30}-\d', str(guide)):
        raise ValueError('Número da guia inválido')
    try:
        datetime.strptime(due, '%d/%m/%Y')
    except ValueError as exc:
        raise ValueError('Vencimento deve ser DD/MM/AAAA') from exc
    path = Path(path)
    if not path.exists() or path.stat().st_size < 100 or path.read_bytes()[:5] != b'%PDF-':
        raise ValueError('Download não é um PDF válido')
    reader = PdfReader(path)
    text = '\n'.join(page.extract_text() or '' for page in reader.pages)
    compact = digits(text)
    employer = re.search(r'CPF/CNPJ do Empregador\s*([\d./-]+)', text)
    employer_cnpj = digits(employer.group(1)) if employer else ''
    if employer_cnpj not in (company['cnpj'], company['cnpj'][:8]):
        raise ValueError('CNPJ do PDF não corresponde à empresa')
    if digits(guide) not in compact:
        raise ValueError('Número da guia não foi confirmado no PDF')
    if due not in text or reais(company['total']) not in text:
        raise ValueError('Vencimento ou total não confirmado no PDF')
    for label, key in [('Total FGTS', 'fgts'), ('Total Consignado', 'consignado')]:
        found = re.search(re.escape(label) + r':\s*([\d.]+,\d{2})', text)
        if not found or money(found.group(1)) != company[key]:
            raise ValueError(f'{label} do PDF não corresponde à empresa')
    # Range guides may abbreviate competence; require both ends explicitly.
    if initial not in text or final not in text:
        raise ValueError('Competência não confirmada no PDF; revisar manualmente')
    return True
